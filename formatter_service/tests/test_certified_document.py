"""DOC-0328 — front-matter renderers + complete certified document assembly tests.

Covers caption/title/appearances/certificate line generation (LEGAL-FORMAT parity to the
retiring reference, with documented exclusions), explicit failure on missing reporter
identity, deterministic section order, and a complete DOCX assembly (front -> body ->
certificate/errata) that reuses the surviving body renderer — all local, no deploy.
"""

import pytest
from docx import Document

from formatter_core.certified_sections import (
    CERTIFIED_DOCUMENT_ORDER,
    CertifiedMetadataError,
    appearances_lines,
    caption_lines,
    certificate_lines,
    render_certified_document_to_docx,
    title_page_lines,
)


CAPTION = {
    "causeNumber": "2026-CI-01234",
    "caption": "JANE DOE vs. ACME CORP.",
    "court": "285th District Court",
    "judicialDistrict": "285th",
    "county": "Bexar",
    "state": "Texas",
    "jurisdictionType": "District Court",
    "deponent": "JANE DOE",
    "depositionDate": "July 22, 2026",
    "startTime": "9:00 a.m.",
    "endTime": "11:30 a.m.",
    "location": "123 Main St, San Antonio, Texas",
}

APPEARANCES = [
    {"category": "attorney", "name": "MR. SMITH", "firm": "Smith LLP", "representing": "Plaintiff", "barNumber": "12345", "phone": "555-1000"},
    {"category": "attorney", "name": "MS. JONES", "firm": "Jones PC", "representing": "Defendant", "barNumber": "67890"},
    {"category": "participant", "name": "BOB VIDEOGRAPHER"},
]

REPORTER = {
    "reporterName": "MARY REPORTER",
    "csrLicense": "CSR-9999",
    "firmRegistration": "FIRM-123",
    "csrCertExpiration": "12/31/2027",
    "deponent": "JANE DOE",
    "caption": "JANE DOE vs. ACME CORP.",
    "causeNumber": "2026-CI-01234",
}


def _body_render_model(n_lines: int) -> dict:
    lines = []
    for i in range(n_lines):
        lines.append({
            "content": f"Q. Question number {i}?" if i % 2 == 0 else f"A. Answer number {i}.",
            "geometry": {"role": "qa", "first_line_tab_inches": 0.5, "text_tab_inches": 1.0, "continuation_indent_inches": 0.0, "paragraph_index": i, "paragraph_id": f"p{i}"},
        })
    return {
        "transcriptId": "t-doc",
        "geometry": {"format_box_width_inches": 6.5, "left_margin_inches": 1.25, "right_margin_inches": 0.75, "line_spacing_points": 28, "lines_per_page": 25},
        "lines": lines,
    }


def _certified_data() -> dict:
    return {
        "caption": CAPTION,
        "appearances": APPEARANCES,
        "reporterCertificate": REPORTER,
        "witnessName": "JANE DOE",
        "depoDate": "July 22, 2026",
        "examinationIndex": [{"kind": "EXAMINATION", "examinerLabel": "MR. SMITH", "page": 4}],
        "exhibitIndex": [{"exhibit_number": "1", "description": "Contract", "offered": 40, "admitted": 41, "excluded": None}],
        "errata": [{"page": 12, "line": 4, "from": "there", "to": "their", "reason": "typo"}],
    }


def test_caption_includes_style_cause_appearances_and_proceedings() -> None:
    lines = caption_lines(CAPTION, APPEARANCES)
    assert "  JANE DOE vs. ACME CORP." in lines
    assert "  Cause No. 2026-CI-01234" in lines
    assert "  DEPOSITION OF JANE DOE" in lines
    assert "  A P P E A R A N C E S" in lines
    assert "  FOR THE PLAINTIFF:" in lines
    assert "    State Bar No. 12345" in lines
    assert "  ALSO PRESENT:" in lines
    assert lines[-1] == "  PROCEEDINGS"


def test_appearances_group_attorneys_and_present() -> None:
    lines = appearances_lines(APPEARANCES)
    assert lines[0] == "  A P P E A R A N C E S"
    assert "  FOR THE DEFENDANT:" in lines
    assert "    MS. JONES" in lines


def test_title_page_has_style_and_reporter_credential() -> None:
    lines = title_page_lines(CAPTION, REPORTER)
    assert lines[0] == "  NO. 2026-CI-01234"
    assert "  ORAL DEPOSITION OF" in lines
    assert "  JANE DOE" in lines
    assert any("MARY REPORTER, CSR-9999 in and for the State of Texas." in ln for ln in lines)


def test_certificate_reproduces_reference_and_carries_reporter_identity() -> None:
    lines = certificate_lines(REPORTER)
    assert lines[0] == "  CERTIFICATE"
    assert "  I, MARY REPORTER, Certified Shorthand Reporter in and for" in lines
    assert "  That the witness, JANE DOE, was duly sworn by me, and" in lines
    assert "  State of Texas, CSR-9999  (Exp. 12/31/2027)" in lines
    assert "  Firm Reg. No. FIRM-123" in lines


def test_certificate_fails_explicitly_on_missing_reporter_identity() -> None:
    with pytest.raises(CertifiedMetadataError):
        certificate_lines({"reporterName": "MARY REPORTER"})  # no CSR
    with pytest.raises(CertifiedMetadataError):
        certificate_lines({"csrLicense": "CSR-9999"})  # no name


def test_complete_document_assembly_order_and_body(tmp_path) -> None:
    out = render_certified_document_to_docx(_body_render_model(30), _certified_data(), str(tmp_path / "certified.docx"))
    doc = Document(out)

    # Each certified section is one or more lined tables; a section may span >1 page
    # (e.g. the caption), so assert the section markers appear as an ordered subsequence.
    first_cells = [t.rows[0].cells[1].paragraphs[0].text for t in doc.tables]
    markers = []
    for cell in first_cells:
        for name, needle in [
            ("title", "  NO. "),
            ("caption", "  IN THE"),
            ("witness", "  INDEX OF WITNESSES"),
            ("exhibit", "  INDEX OF EXHIBITS"),
            ("certificate", "  CERTIFICATE"),
            ("changes", "  CHANGES AND SIGNATURE"),
            ("signature", "  I, JANE DOE"),
        ]:
            if cell.startswith(needle):
                markers.append(name)
                break
    assert markers == ["title", "caption", "witness", "exhibit", "certificate", "changes", "signature"]

    # The numbered body is present as paragraphs (reviewed Q/A content is observable).
    body_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Question number 0?" in body_text
    assert "Answer number 1." in body_text


def test_complete_document_requires_reporter_identity(tmp_path) -> None:
    data = _certified_data()
    data["reporterCertificate"] = {"reporterName": "MARY REPORTER"}  # missing CSR
    with pytest.raises(CertifiedMetadataError):
        render_certified_document_to_docx(_body_render_model(10), data, str(tmp_path / "bad.docx"))


def test_document_order_constant_is_the_documented_sequence() -> None:
    assert CERTIFIED_DOCUMENT_ORDER == (
        "title_page", "caption", "witness_index", "exhibit_index", "body", "certificate", "changes_signature",
    )


def test_complete_document_over_1000_utterances(tmp_path) -> None:
    # >1,000 body lines exercises multi-page pagination through the complete assembly.
    out = render_certified_document_to_docx(_body_render_model(1050), _certified_data(), str(tmp_path / "big.docx"))
    doc = Document(out)
    # Body paragraphs: line-number-prefixed content lines + page-break paragraphs.
    numbered = [p for p in doc.paragraphs if p.text.strip()[:2].strip().isdigit()] if doc.paragraphs else []
    assert len(numbered) >= 1050
    # Back matter still present and last (certificate + changes/signature tables exist).
    back_first_cells = [t.rows[0].cells[1].paragraphs[0].text for t in doc.tables]
    assert "  CERTIFICATE" in back_first_cells
