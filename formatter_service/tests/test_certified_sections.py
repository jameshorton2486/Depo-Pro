"""DOC-0328 — certified back-matter renderer tests (format parity + observable DOCX).

Asserts the re-homed formatter_core certified sections reproduce the reference
transcript_formatter/spec_engine/pages FORMAT (headers, column widths, jurat
boilerplate) while being driven by the derived certified data, in the correct certified
order, and that they render to a real DOCX locally.
"""

from docx import Document

from formatter_core.certified_sections import (
    LINES_PER_PAGE,
    build_certified_back_matter_lines,
    changes_page_lines,
    examination_index_lines,
    exhibit_index_lines,
    render_certified_back_matter_to_docx,
    signature_page_lines,
)


def test_examination_index_reproduces_reference_header_and_pivots_kinds() -> None:
    lines = examination_index_lines(
        "JANE DOE",
        [
            {"kind": "EXAMINATION", "examinerLabel": "MR. SMITH", "page": 4},
            {"kind": "CROSS-EXAMINATION", "examinerLabel": "MS. JONES", "page": 30},
            {"kind": "REDIRECT", "examinerLabel": "MR. SMITH", "page": 45},
        ],
    )
    assert lines[0] == "  INDEX OF WITNESSES"
    # Column header byte-identical to the reference witness_index.py.
    assert lines[2] == f"  {'WITNESS':<30}{'DIR':>5}{'CRS':>5}{'REDIR':>6}{'RECRSS':>7}{'VD':>4}"
    # The flat examinations pivot into DIR/CRS/REDIR columns; RECRSS/VD stay blank.
    assert lines[4] == f"  {'JANE DOE':<30}{'4':>5}{'30':>5}{'45':>6}{'':>7}{'':>4}"


def test_exhibit_index_reproduces_reference_format_with_derived_pages() -> None:
    lines = exhibit_index_lines(
        [
            {"exhibit_number": "1", "description": "Employment agreement", "offered": 40, "admitted": 41, "excluded": None},
            {"exhibit_number": "2", "marked": 12, "offered": None, "admitted": None, "excluded": 55},
        ]
    )
    assert lines[0] == "  INDEX OF EXHIBITS"
    assert lines[2] == f"  {'NO.':<6}{'DESCRIPTION':<36}{'OFFERED':>8}{'ADMITTED':>9}{'EXCLUDED':>9}"
    assert lines[4] == f"  {'1':<6}{'Employment agreement':<36}{'40':>8}{'41':>9}{'':>9}"
    # Absent actions render blank, never a fabricated page.
    assert lines[5] == f"  {'2':<6}{'':<36}{'':>8}{'':>9}{'55':>9}"


def test_changes_page_grid_and_padding() -> None:
    lines = changes_page_lines(
        "JANE DOE",
        "2026-07-22",
        [{"page": 12, "line": 4, "from": "there", "to": "their", "reason": "typo"}],
    )
    assert lines[0] == "  CHANGES AND SIGNATURE"
    assert lines[2] == f"  WITNESS NAME: {'JANE DOE':<28}  DATE: 2026-07-22"
    assert lines[3] == f"  {'PAGE':<8}{'LINE':<8}{'CHANGE':<30}REASON"
    # First errata row cites the certified page/line from the derived data.
    assert lines[4] == f"  {'12':<8}{'4':<8}{'there -> their':<30}typo"
    # Remaining 19 rows are blank fill; page padded to exactly 25 lines.
    assert lines[5].startswith("  ______")
    assert len(lines) == LINES_PER_PAGE


def test_signature_page_jurat_boilerplate_is_verbatim() -> None:
    lines = signature_page_lines("Jane Doe", notary_county="Bexar", notary_name="A. Notary")
    assert lines[0] == "  I, Jane Doe, have read the foregoing"
    assert lines[6] == "  (JANE DOE)"
    assert lines[8] == "  THE STATE OF TEXAS              )"
    assert lines[9] == f"  COUNTY OF {'Bexar'}    )"
    assert lines[10] == "  Before me, A. Notary, on this day personally"
    assert "  NOTARY PUBLIC IN AND FOR" in lines
    assert "  THE STATE OF TEXAS" in lines


def test_back_matter_assembled_in_certified_order() -> None:
    pages = build_certified_back_matter_lines(
        {
            "witnessName": "JANE DOE",
            "depoDate": "2026-07-22",
            "examinationIndex": [{"kind": "EXAMINATION", "examinerLabel": "MR. SMITH", "page": 4}],
            "exhibitIndex": [{"exhibit_number": "1", "offered": 40, "admitted": 41, "excluded": None}],
            "errata": [{"page": 12, "line": 4, "from": "there", "to": "their", "reason": "typo"}],
        }
    )
    firsts = [page[0] for page in pages]
    assert firsts == [
        "  INDEX OF WITNESSES",
        "  INDEX OF EXHIBITS",
        "  CHANGES AND SIGNATURE",
        "  I, JANE DOE, have read the foregoing",
    ]


def test_empty_indexes_still_produce_changes_and_signature() -> None:
    pages = build_certified_back_matter_lines({"witnessName": "JANE DOE", "errata": []})
    firsts = [page[0] for page in pages]
    # No index pages when there are no examinations/exhibits; changes + signature always present.
    assert firsts == ["  CHANGES AND SIGNATURE", "  I, JANE DOE, have read the foregoing"]


def test_renders_observable_docx(tmp_path) -> None:
    out = render_certified_back_matter_to_docx(
        {
            "witnessName": "JANE DOE",
            "depoDate": "2026-07-22",
            "examinationIndex": [{"kind": "EXAMINATION", "examinerLabel": "MR. SMITH", "page": 4}],
            "exhibitIndex": [{"exhibit_number": "1", "offered": 40, "admitted": 41, "excluded": None}],
            "errata": [{"page": 12, "line": 4, "from": "there", "to": "their", "reason": "typo"}],
        },
        str(tmp_path / "back_matter.docx"),
    )
    doc = Document(out)
    # Each certified page is one 25-row lined table: witness index, exhibit index,
    # changes grid, signature/notary => 4 tables.
    assert len(doc.tables) == 4
    for table in doc.tables:
        assert len(table.rows) == LINES_PER_PAGE
        assert len(table.columns) == 2
