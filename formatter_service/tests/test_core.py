from pathlib import Path

from docx import Document

from formatter_core import format_render_model


def render_model(lines: list[dict[str, object]]) -> dict[str, object]:
    return {
        "transcriptId": "transcript-001",
        "geometry": {
            "format_box_width_inches": 6.5,
            "left_margin_inches": 1.25,
            "right_margin_inches": 0.75,
            "line_spacing_points": 28,
            "lines_per_page": 25,
        },
        "lines": lines,
    }


def line(content: str, role: str, first_tab: float, text_tab: float | None = None, continuation: float = 0) -> dict[str, object]:
    return {
        "content": content,
        "geometry": {
            "role": role,
            "first_line_tab_inches": first_tab,
            "text_tab_inches": text_tab,
            "continuation_indent_inches": continuation,
        },
    }


def test_generates_docx_from_completed_render_model(tmp_path: Path) -> None:
    artifacts = format_render_model(
        render_model([line("Q. Where were you?", "qa", 0.5, 1.0, 1.0)]),
        ["DOCX"],
        tmp_path,
    )

    assert artifacts["DOCX"].suffix == ".docx"
    assert artifacts["DOCX"].exists()


def test_preserves_geometry_without_numbered_blank_lines(tmp_path: Path) -> None:
    path = format_render_model(
        render_model([
            line("Q. Where were you?", "qa", 0.5, 1.0, 1.0),
            line("A. At the synthetic office.", "qa", 0.5, 1.0, 1.0),
            line("(Discussion off the record.)", "parenthetical", 2.0),
        ]),
        ["DOCX"],
        tmp_path,
    )["DOCX"]

    paragraphs = Document(path).paragraphs
    assert len(paragraphs) == 3
    assert [paragraph.text.strip() for paragraph in paragraphs] == [
        "1 \tQ.\tWhere were you?",
        "2 \tA.\tAt the synthetic office.",
        "3 \t(Discussion off the record.)",
    ]
    assert "720" in paragraphs[0]._p.xml
    assert "1440" in paragraphs[0]._p.xml
    assert "2880" in paragraphs[2]._p.xml


def test_paginates_wrapped_physical_lines(tmp_path: Path) -> None:
    long_answer = "A. " + "word " * 170
    path = format_render_model(
        render_model([line(long_answer, "qa", 0.5, 1.0, 1.0)]),
        ["DOCX"],
        tmp_path,
    )["DOCX"]

    document = Document(path)
    transcript_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    assert len(transcript_paragraphs) > 1
    assert transcript_paragraphs[0].text.lstrip().startswith("1 \tA.\t")
    assert transcript_paragraphs[1].text.lstrip().startswith("2 \t")

def test_consumes_all_geometry_without_inferring_role_from_text(tmp_path: Path) -> None:
    model = render_model([
        line("Q. " + "synthetic " * 80, "speaker", 1.75, None, 0.25),
    ])
    geometry = model["geometry"]
    assert isinstance(geometry, dict)
    geometry["line_spacing_points"] = 30
    geometry["left_margin_inches"] = 1.4
    geometry["right_margin_inches"] = 0.6
    geometry["lines_per_page"] = 20

    path = format_render_model(model, ["DOCX"], tmp_path)["DOCX"]
    document = Document(path)
    first, continuation = document.paragraphs[0:2]

    assert round(document.sections[0].left_margin.inches, 2) == 1.4
    assert round(document.sections[0].right_margin.inches, 2) == 0.6
    assert first.paragraph_format.line_spacing.pt == 30
    assert "w:pos=\"2520\"" in first._p.xml
    assert "w:pos=\"720\"" not in first._p.xml
    assert "w:pos=\"1440\"" not in first._p.xml
    assert "w:pos=\"360\"" in continuation._p.xml
    assert first.text.lstrip().startswith("1 \tQ.")
