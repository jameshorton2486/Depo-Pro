"""DOC-0328 — cross-runtime output proof, Python renderer half.

Consumes the committed certified-transport fixture emitted by the TS producer
(src/lib/export/certifiedTransport.test.ts) and renders it to a complete certified DOCX,
then asserts semantically significant content from BOTH the reviewed transcript body and
the certified sections. Together with the TS half this closes:

  persisted reviewed line_type -> Working Transcript -> FinalizedTranscriptModel
    -> serialized certified transport (fixture) -> formatter_core -> complete DOCX.

The fixture is the seam: the TS test proves the pipeline emits exactly this contract; this
test proves the same contract renders to a correct certified document.
"""

import json
from pathlib import Path

from docx import Document

from formatter_core.certified_sections import render_certified_document_to_docx

# Mirrors CERTIFIED_TRANSPORT_VERSION in src/lib/export/certifiedTransport.ts.
EXPECTED_TRANSPORT_VERSION = "2026-08-11"

_FIXTURE = (
    Path(__file__).resolve().parents[2]
    / "src" / "lib" / "export" / "__fixtures__" / "certifiedTransport.fixture.json"
)


def _load_fixture() -> dict:
    with _FIXTURE.open(encoding="utf-8") as handle:
        return json.load(handle)


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_fixture_is_the_pinned_transport_version() -> None:
    payload = _load_fixture()
    assert payload["version"] == EXPECTED_TRANSPORT_VERSION


def test_cross_runtime_render_carries_reviewed_body_and_certified_sections(tmp_path) -> None:
    payload = _load_fixture()
    out = render_certified_document_to_docx(
        payload["renderModel"], payload["certified"], str(tmp_path / "certified.docx")
    )
    doc = Document(out)
    text = _all_text(doc)

    # (1) Reviewed line_type reached observable output: the qa_split produced distinct Q
    # and A units, both present in the rendered body.
    assert "You understand the oath?" in text
    assert "Yes I understand completely." in text

    # (2) Certified front matter from the transported metadata.
    assert "JANE DOE" in text                 # witness / deponent
    assert "2026-CI-01234" in text            # cause number
    assert "A P P E A R A N C E S" in text
    assert "MR. SMITH" in text                # appearance

    # (3) Reporter certificate (required certification facts present, not fabricated).
    assert "CERTIFICATE" in text
    assert "MARY REPORTER" in text
    assert "CSR-9999" in text

    # (4) Exhibit index derived from the body parenthetical.
    assert "INDEX OF EXHIBITS" in text

    # (5) Errata cites the certified page/line resolved over the canonical PaginationMap.
    errata = payload["certified"]["errata"][0]
    assert "CHANGES AND SIGNATURE" in text
    assert f"completely -> fully" in text
    assert str(errata["page"]) in text and str(errata["line"]) in text


def test_cross_runtime_document_section_order(tmp_path) -> None:
    payload = _load_fixture()
    out = render_certified_document_to_docx(
        payload["renderModel"], payload["certified"], str(tmp_path / "certified.docx")
    )
    doc = Document(out)
    firsts = [t.rows[0].cells[1].paragraphs[0].text for t in doc.tables]
    markers = [c for c in firsts if c.startswith("  NO. ") or c.startswith("  IN THE") or c in ("  INDEX OF EXHIBITS", "  CERTIFICATE", "  CHANGES AND SIGNATURE")]
    # title -> caption -> exhibit index -> certificate -> changes (witness index absent: no examinations detected)
    assert markers[0].startswith("  NO. ")
    assert markers[-2] == "  CERTIFICATE"
    assert markers[-1] == "  CHANGES AND SIGNATURE"
    assert "  INDEX OF EXHIBITS" in markers
