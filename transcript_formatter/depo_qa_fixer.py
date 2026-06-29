"""
depo_qa_fixer.py — Standalone Q/A split and cleanup tool.

Applies Rules 1-3 and optional scopist flag removal to a formatted
deposition DOCX. Safe to run repeatedly — non-destructive.
"""

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")


ANSWER_TOKENS = [
    'yes', 'no', 'correct', 'right', 'yeah', 'yep', 'yup', 'nope', 'nah',
    'mm-hmm', 'uh-huh', 'no, sir', 'yes, sir', "no, ma'am", "yes, ma'am",
    'i do', 'i do not', "i don't", 'i did', 'i did not',
    'i have', 'i have not', 'i would', 'i could', 'i could not',
    'i was not', 'i am not', 'i was', 'i will', 'i will not',
    'i believe', 'i think', 'i recall', 'i remember',
    "i don't know", "i don't recall", "i don't remember",
    'we have', 'we did', 'we were', 'they would have',
    'there was not', 'there is not', 'there were not',
    "that's correct", 'that is correct', 'not particularly', 'not necessarily',
    'not that i recall', 'sure', 'never', 'absolutely', 'definitely',
    'fentanyl', 'in the ed', 'surgical issues',
]

ANSWER_TOKEN_RE = re.compile(
    r'(\?)\s{1,2}'
    r'(No\.|Yes\.|Correct\.|Right\.|Yeah\.|Yep\.|Yup\.|Nope\.|Nah\.'
    r'|Mm-hmm\.|Uh-huh\.|No,\s+sir\.|Yes,\s+sir\.|No,\s+ma\'am\.|Yes,\s+ma\'am\.'
    r'|I have\.|I have not\.|I do\.|I do not\.|I did\.|I did not\.'
    r'|I would\.|I could\.|I could not\.|I was not\.|I am not\.'
    r'|I recall\.|I remember\.|Sure\.|Never\.|That\'s correct\.|That is correct\.'
    r'|Not particularly\.|Not necessarily\.|Not that I recall\.'
    r'|We have\.|We did\.|We were\.|There was not\.|There is not\.'
    r'|They would have\.)',
    re.IGNORECASE
)

CORRECT_MID_RE = re.compile(
    r'(\.\s{1,2})(Correct\.)(\s{1,2}[A-Z])',
    re.IGNORECASE
)

TRAILING_OKAY_RE = re.compile(r'\s{1,2}Okay\.$')
SCOPIST_FLAG_RE = re.compile(r'\[SCOPIST:.*?\]', re.DOTALL)


def split_on_answer_token(q_text: str):
    match = ANSWER_TOKEN_RE.search(q_text)
    if not match:
        return None
    question_part = q_text[:match.start(2)].strip()
    rest = q_text[match.start(2):].strip()
    end_match = re.search(r'(?<=[.!?])\s+[A-Z]', rest)
    if end_match:
        answer_part = rest[:end_match.start()].strip()
        continuation = rest[end_match.start():].strip()
        return [question_part, answer_part, continuation]
    return [question_part, rest]


def split_correct_mid(q_text: str):
    match = CORRECT_MID_RE.search(q_text)
    if not match:
        return None
    before = q_text[:match.start(2)].rstrip()
    continuation = q_text[match.end(2):].strip()
    return (before, 'Correct.', continuation)


def fix_trailing_okay(paragraphs: list) -> list:
    result = list(paragraphs)
    i = 0
    while i < len(result):
        p = result[i]
        if p['label'] == 'A' and TRAILING_OKAY_RE.search(p['text']):
            clean_text = TRAILING_OKAY_RE.sub('', p['text']).strip()
            result[i] = {**p, 'text': clean_text}
            for j in range(i + 1, len(result)):
                if result[j]['label'] == 'Q':
                    next_q = result[j]
                    if not next_q['text'].startswith('Okay.'):
                        result[j] = {**next_q, 'text': 'Okay. ' + next_q['text']}
                    break
        i += 1
    return result


def remove_scopist_flags(text: str) -> str:
    """Strip only the bracketed [SCOPIST: FLAG ...] span and keep the testimony token."""
    return SCOPIST_FLAG_RE.sub('', text).strip()


def get_paragraph_label(para) -> str:
    text = para.text.strip()
    if text.startswith('Q.') or text.startswith('Q '):
        return 'Q'
    if text.startswith('A.') or text.startswith('A '):
        return 'A'
    return 'OTHER'


def process_docx(input_path: Path, output_path: Path, clean: bool = False):
    doc = Document(str(input_path))
    paragraphs = [
        {'label': get_paragraph_label(p), 'text': p.text.strip(), 'para': p}
        for p in doc.paragraphs
    ]

    changes = 0
    paragraphs = fix_trailing_okay(paragraphs)
    new_paragraphs = []
    for p in paragraphs:
        if p['label'] != 'Q':
            new_paragraphs.append(p)
            continue

        text = p['text']
        split2 = split_correct_mid(text)
        if split2:
            new_paragraphs.append({'label': 'Q', 'text': split2[0], 'para': p['para']})
            new_paragraphs.append({'label': 'A', 'text': split2[1], 'para': None})
            if len(split2) > 2 and split2[2]:
                new_paragraphs.append({'label': 'Q', 'text': split2[2], 'para': None})
            changes += 1
            continue

        split1 = split_on_answer_token(text)
        if split1:
            new_paragraphs.append({'label': 'Q', 'text': split1[0], 'para': p['para']})
            new_paragraphs.append({'label': 'A', 'text': split1[1], 'para': None})
            if len(split1) > 2 and split1[2]:
                new_paragraphs.append({'label': 'Q', 'text': split1[2], 'para': None})
            changes += 1
            continue

        new_paragraphs.append(p)

    new_doc = Document()
    for p in new_paragraphs:
        text = p['text']
        if clean:
            text = remove_scopist_flags(text)
        if text:
            new_doc.add_paragraph(text)

    new_doc.save(str(output_path))
    return changes


def main():
    parser = argparse.ArgumentParser(description='Fix Q/A splits in deposition DOCX.')
    parser.add_argument('input', help='Input DOCX file')
    parser.add_argument('--out', help='Output filename (default: input_FIXED.docx)')
    parser.add_argument('--clean', action='store_true', help='Remove scopist flags for clean delivery version')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        sys.exit(f"File not found: {input_path}")

    output_path = Path(args.out) if args.out else input_path.with_stem(input_path.stem + '_FIXED')
    changes = process_docx(input_path, output_path, clean=args.clean)
    print(f"Saved: {output_path.name}")
    print(f"Total changes: {changes}")


if __name__ == '__main__':
    main()
