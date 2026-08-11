"""certified_sections.py — DOC-0328 certified back-matter renderer (surviving home).

Reproduces the fixed-width certified section FORMAT of the retiring
transcript_formatter/spec_engine/pages builders (INDEX OF WITNESSES, INDEX OF EXHIBITS,
CHANGES AND SIGNATURE + notary jurat), but driven by the DERIVED certified data — the
examination/exhibit index rows and errata resolved over the canonical PaginationMap —
instead of externally-supplied page-ref strings.

The legal/typographic FORMAT is preserved verbatim (headers, column widths, separators,
jurat boilerplate) so the re-homed output matches the reference; only the data SOURCE
changes (canonical pagination, one authority — no second paginator here). Page numbers
arrive already computed.

Pure line generators + a 25-row lined-page DOCX writer, so the sections are observable
as real DOCX locally without any deployment.
"""

from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# ── lined-page geometry (mirrors the reference _lined_page primitive) ──────────
_FONT = "Courier New"
_FONT_SIZE = Pt(12)
_COLOR_GRAY = RGBColor(0x99, 0x99, 0x99)
_COLOR_BLK = RGBColor(0x00, 0x00, 0x00)
LINES_PER_PAGE = 25
_NUM_COL_TWIPS = 540
_TOTAL_TWIPS = 9000
_CONTENT_TWIPS = _TOTAL_TWIPS - _NUM_COL_TWIPS
_ROW_HEIGHT_TWIPS = 360

_SEP = "─" * 60
CHANGE_ROWS_PAGE1 = 20

# Examination kind -> witness-index column. Mirrors the reference DIR/CRS/REDIR/RECRSS/VD
# columns; the finalized examination index is a flat list, so it is pivoted here.
_KIND_TO_COLUMN = {
    "EXAMINATION": "direct",
    "CROSS-EXAMINATION": "cross",
    "REDIRECT": "redirect",
    "RECROSS": "recross",
    "VOIR_DIRE": "voir_dire",
}


def _page_str(value: object) -> str:
    """A page number as its printed string, or "" for a null/absent page."""
    if value is None:
        return ""
    return str(value)


def examination_index_lines(
    witness_name: str | None,
    examinations: Sequence[Mapping[str, object]],
) -> list[str]:
    """INDEX OF WITNESSES content lines, pivoting the flat examination index into the
    reference DIR/CRS/REDIR/RECRSS/VD columns (first occurrence of each kind wins)."""
    columns = {"direct": "", "cross": "", "redirect": "", "recross": "", "voir_dire": ""}
    for exam in examinations:
        column = _KIND_TO_COLUMN.get(str(exam.get("kind")))
        if column and not columns[column]:
            columns[column] = _page_str(exam.get("page"))

    name = witness_name or "[WITNESS NAME]"
    return [
        "  INDEX OF WITNESSES",
        "",
        f"  {'WITNESS':<30}{'DIR':>5}{'CRS':>5}{'REDIR':>6}{'RECRSS':>7}{'VD':>4}",
        f"  {_SEP}",
        f"  {name:<30}"
        f"{columns['direct']:>5}"
        f"{columns['cross']:>5}"
        f"{columns['redirect']:>6}"
        f"{columns['recross']:>7}"
        f"{columns['voir_dire']:>4}",
    ]


def exhibit_index_lines(exhibits: Sequence[Mapping[str, object]]) -> list[str]:
    """INDEX OF EXHIBITS content lines from the finalized exhibit index rows. The
    description column is filled from an optional `description` on each row (sourced from
    the CaseRecord upstream); page columns are the derived offered/admitted/excluded."""
    header = [
        "  INDEX OF EXHIBITS",
        "",
        f"  {'NO.':<6}{'DESCRIPTION':<36}{'OFFERED':>8}{'ADMITTED':>9}{'EXCLUDED':>9}",
        f"  {_SEP}",
    ]
    body: list[str] = []
    for exhibit in exhibits:
        description = str(exhibit.get("description") or "")[:35]
        body.append(
            f"  {str(exhibit.get('exhibit_number', '')):<6}{description:<36}"
            f"{_page_str(exhibit.get('offered')):>8}"
            f"{_page_str(exhibit.get('admitted')):>9}"
            f"{_page_str(exhibit.get('excluded')):>9}"
        )
    return header + body


def changes_page_lines(
    witness_name: str | None,
    depo_date: str | None,
    errata: Sequence[Mapping[str, object]],
) -> list[str]:
    """Page 1 of CHANGES AND SIGNATURE: the 20-row errata grid, padded to 25 lines. Each
    errata row cites the certified page/line resolved over the canonical PaginationMap;
    the change is rendered as "from -> to"."""
    name = witness_name or ""
    lines = [
        "  CHANGES AND SIGNATURE",
        "",
        f"  WITNESS NAME: {name:<28}  DATE: {depo_date or ''}",
        f"  {'PAGE':<8}{'LINE':<8}{'CHANGE':<30}REASON",
    ]
    for idx in range(CHANGE_ROWS_PAGE1):
        if idx < len(errata):
            row = errata[idx]
            change_text = f"{row.get('from', '')} -> {row.get('to', '')}"[:29]
            lines.append(
                f"  {_page_str(row.get('page')):<8}{_page_str(row.get('line')):<8}{change_text:<30}{row.get('reason', '')}"
            )
        else:
            lines.append(
                f"  {'______':<8}{'______':<8}{'______________________':<30}__________________"
            )
    while len(lines) < LINES_PER_PAGE:
        lines.append("")
    return lines


def signature_page_lines(
    witness_name: str | None,
    notary_county: str | None = None,
    notary_name: str | None = None,
    identification_method: str | None = None,
) -> list[str]:
    """Page 2 of CHANGES AND SIGNATURE: the signature block + notary jurat. Boilerplate
    text is reproduced verbatim from the reference; only witness/notary fields vary."""
    name = witness_name or ""
    id_method = identification_method or (
        "driver's license or other government-issued photo identification"
    )
    return [
        f"  I, {name}, have read the foregoing",
        "  deposition and hereby affix my signature that same is",
        "  true and correct, except as noted above.",
        "",
        "",
        "  ______________________________",
        f"  ({name.upper()})",
        "",
        "  THE STATE OF TEXAS              )",
        f"  COUNTY OF {notary_county or '___________________'}    )",
        f"  Before me, {notary_name or '________________'}, on this day personally",
        f"  appeared {name} known to me (or",
        f"  proved to me under oath or through {id_method})",
        "  (description of identity card or other document) to be the",
        "  person whose name is subscribed to the foregoing instrument",
        "  and acknowledged to me that they executed the same for the",
        "  purposes and consideration therein expressed.",
        "  Given under my hand and seal of office this _____",
        "  day of _____________________, _______.",
        "",
        "",
        "",
        "  NOTARY PUBLIC IN AND FOR",
        "  THE STATE OF TEXAS",
        "",
    ]


class CertifiedMetadataError(ValueError):
    """Raised when metadata required to certify a page is absent. The certificate must
    fail explicitly rather than emit a fabricated reporter identity / CSR / signature."""


_RULE = "─" * 60


def _appearance_group_lines(appearances: Sequence[Mapping[str, object]]) -> list[str]:
    """Group appearances the way the reference caption does: attorneys under
    'FOR THE <party>:' by who they represent, everyone else under 'ALSO PRESENT:'."""
    attorneys = [a for a in appearances if str(a.get("category") or "") == "attorney"]
    others = [a for a in appearances if str(a.get("category") or "") != "attorney"]

    lines: list[str] = []
    order: list[str] = []
    by_party: dict[str, list[Mapping[str, object]]] = {}
    for attorney in attorneys:
        party = str(attorney.get("representing") or "").strip() or "THE RECORD"
        if party not in by_party:
            by_party[party] = []
            order.append(party)
        by_party[party].append(attorney)

    for party in order:
        lines.append(f"  FOR THE {party.upper()}:")
        for attorney in by_party[party]:
            lines.append(f"    {attorney.get('name') or ''}")
            if attorney.get("barNumber"):
                lines.append(f"    State Bar No. {attorney.get('barNumber')}")
            if attorney.get("firm"):
                lines.append(f"    {attorney.get('firm')}")
            if attorney.get("phone"):
                lines.append(f"    {attorney.get('phone')}")
            lines.append("")

    present = [a for a in others if a.get("name")]
    if present:
        lines.append("  ALSO PRESENT:")
        for entry in present:
            lines.append(f"    {entry.get('name')}")
        lines.append("")

    return lines


def appearances_lines(appearances: Sequence[Mapping[str, object]]) -> list[str]:
    """Standalone A P P E A R A N C E S section content."""
    return ["  A P P E A R A N C E S", ""] + _appearance_group_lines(appearances)


def caption_lines(
    caption: Mapping[str, object],
    appearances: Sequence[Mapping[str, object]],
) -> list[str]:
    """Caption page content (Texas district-court style), LEGAL-FORMAT parity to the
    reference caption.py. The canonical envelope carries a single case-style string plus
    separate parties (not the reference's plaintiff_name/defendant_names[] split), and
    lacks court_type / witness_title / subpoena_duces_tecum — those reference elements
    are intentionally excluded rather than fabricated."""
    court_heading = str(caption.get("court") or caption.get("jurisdictionType") or "").upper()
    witness = str(caption.get("deponent") or "[WITNESS NAME]").upper()

    lines = [
        f"  IN THE {court_heading}".rstrip(),
    ]
    if caption.get("judicialDistrict"):
        lines.append(f"  {str(caption.get('judicialDistrict')).upper()}")
    if caption.get("county"):
        lines.append(f"  {str(caption.get('county')).upper()} COUNTY, TEXAS")
    lines += [
        f"  {_RULE}",
        f"  {caption.get('caption') or '[CASE STYLE]'}",
        f"  Cause No. {caption.get('causeNumber') or '____________'}",
        f"  {_RULE}",
        f"  DEPOSITION OF {witness}",
    ]
    if caption.get("depositionDate"):
        lines.append(f"  Taken on {caption.get('depositionDate')}")
    lines += [f"  {_RULE}", "  A P P E A R A N C E S", ""]
    lines += _appearance_group_lines(appearances)
    lines += [f"  {_RULE}", "  PROCEEDINGS"]
    return lines


def title_page_lines(
    caption: Mapping[str, object],
    reporter: Mapping[str, object],
) -> list[str]:
    """Style/Title page content (UFM Fig03), LEGAL-FORMAT parity to title_page.py. Uses
    the case-style string in place of the reference plaintiff/defendant split; the
    videotaped flag is unknown in the canonical envelope so the neutral 'ORAL' label is
    used (documented exclusion)."""
    witness = str(caption.get("deponent") or "[WITNESS NAME]")
    lines = [
        f"  NO. {caption.get('causeNumber') or '____________'}",
        "",
        f"  {caption.get('caption') or '[CASE STYLE]'}",
    ]
    if caption.get("county"):
        lines.append(f"  {str(caption.get('county')).upper()} COUNTY, TEXAS")
    if caption.get("judicialDistrict"):
        lines.append(f"  {caption.get('judicialDistrict')} JUDICIAL DISTRICT")
    lines += [
        f"  {_RULE}",
        "  ORAL DEPOSITION OF",
        f"  {witness.upper()}",
        f"  {caption.get('depositionDate') or ''}".rstrip(),
        f"  {_RULE}",
        "",
        f"  ORAL DEPOSITION OF {witness}, produced as a witness,",
        "  duly sworn, was taken in the above-styled and numbered cause",
    ]
    span = _time_span(caption.get("startTime"), caption.get("endTime"))
    if span:
        lines.append(f"  {span}, before")
    reporter_name = reporter.get("reporterName")
    reporter_csr = reporter.get("csrLicense")
    if reporter_name:
        credential = f"{reporter_name}, {reporter_csr}" if reporter_csr else str(reporter_name)
        lines.append(f"  {credential} in and for the State of Texas.")
    return lines


def _time_span(start: object, end: object) -> str:
    start_s = str(start).strip() if start else ""
    end_s = str(end).strip() if end else ""
    if start_s and end_s:
        return f"on the record from {start_s} to {end_s}"
    return ""


def certificate_lines(
    certificate: Mapping[str, object],
    *,
    reporter_firm: str | None = None,
    reporter_address: str | None = None,
    reporter_phone: str | None = None,
    waiver: bool = True,
) -> list[str]:
    """Reporter's CERTIFICATE content (Texas), reproducing certificate.py. Reporter
    identity and CSR are certification facts: if either is missing this raises
    CertifiedMetadataError rather than emit a blank/fabricated certificate. Reporter
    firm/address/phone are absent from the canonical envelope and are optional here."""
    reporter_name = certificate.get("reporterName")
    reporter_csr = certificate.get("csrLicense")
    if not reporter_name:
        raise CertifiedMetadataError("certificate requires the reporter's name")
    if not reporter_csr:
        raise CertifiedMetadataError("certificate requires the reporter's CSR license number")

    witness = str(certificate.get("deponent") or "the witness")
    parties = str(certificate.get("caption") or "[parties]")
    csr_str = f"State of Texas, {reporter_csr}"
    if certificate.get("csrCertExpiration"):
        csr_str += f"  (Exp. {certificate.get('csrCertExpiration')})"

    lines = [
        "  CERTIFICATE",
        "",
        f"  I, {reporter_name}, Certified Shorthand Reporter in and for",
        "  the State of Texas, do hereby certify:",
        "",
        f"  That the witness, {witness}, was duly sworn by me, and",
        "  that the transcript of the oral deposition is a true record of",
        "  the testimony given by the witness;",
        "",
    ]
    if waiver:
        lines += [
            "  That examination and signature of the witness to the",
            "  deposition transcript was waived by the witness and the",
            "  parties at the time of the deposition;",
            "",
        ]
    lines += [
        "  That the original deposition transcript was delivered in the",
        f"  matter of {parties}.",
        "",
        "  That I am not a relative, employee, attorney, or counsel of",
        "  any of the parties, nor am I a relative or employee of such",
        "  attorney or counsel, nor am I financially interested in the action.",
        "",
        "",
        "  ______________________________",
        f"  {str(reporter_name).upper()}",
        "  Certified Shorthand Reporter",
        f"  {csr_str}",
    ]
    if certificate.get("firmRegistration"):
        lines.append(f"  Firm Reg. No. {certificate.get('firmRegistration')}")
    if reporter_firm:
        lines.append(f"  {reporter_firm}")
    if reporter_address:
        lines.append(f"  {reporter_address}")
    if reporter_phone:
        lines.append(f"  {reporter_phone}")
    return lines


def paginate_lines(lines: Sequence[str]) -> list[list[str]]:
    """Split a flat list of strings into 25-line pages (mirrors the reference)."""
    if not lines:
        return [[]]
    return [list(lines[i : i + LINES_PER_PAGE]) for i in range(0, len(lines), LINES_PER_PAGE)]


def _tbl_border_xml() -> OxmlElement:
    el = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        child = OxmlElement(f"w:{side}")
        child.set(qn("w:val"), "single")
        child.set(qn("w:sz"), "8")
        child.set(qn("w:space"), "0")
        child.set(qn("w:color"), "000000")
        el.append(child)
    return el


def _set_cell_width(cell, twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(twips))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def _set_row_height(row, twips: int) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_h = OxmlElement("w:trHeight")
    tr_h.set(qn("w:val"), str(twips))
    tr_h.set(qn("w:hRule"), "exact")
    tr_pr.append(tr_h)


def _write_cell_text(cell, text: str, *, align: str = "left", color=None) -> None:
    para = cell.paragraphs[0]
    para.clear()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.font.name = _FONT
    run.font.size = _FONT_SIZE
    if color is not None:
        run.font.color.rgb = color


def write_lined_page(doc: Document, content_lines: Sequence[str]) -> None:
    """Write one 25-line page as a bordered two-column (number, content) table. Callers
    must paginate first; content beyond 25 lines is truncated."""
    lines = list(content_lines[:LINES_PER_PAGE])
    while len(lines) < LINES_PER_PAGE:
        lines.append("")

    table = doc.add_table(rows=LINES_PER_PAGE, cols=2)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_pr.append(_tbl_border_xml())
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(_TOTAL_TWIPS))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    for idx, text in enumerate(lines):
        row = table.rows[idx]
        _set_row_height(row, _ROW_HEIGHT_TWIPS)
        left = row.cells[0]
        right = row.cells[1]
        _set_cell_width(left, _NUM_COL_TWIPS)
        _set_cell_width(right, _CONTENT_TWIPS)
        _write_cell_text(left, str(idx + 1), align="right", color=_COLOR_GRAY)
        _write_cell_text(right, text, color=_COLOR_BLK)


def build_certified_back_matter_lines(sections: Mapping[str, object]) -> list[list[str]]:
    """Assemble the certified back-matter as a list of 25-line pages, in certified order:
    witness index, exhibit index, changes grid, signature/notary. Empty sections are
    skipped. `sections` mirrors the TS certifiedSectionModel output plus the notary
    fields; page numbers are already resolved over the canonical PaginationMap."""
    pages: list[list[str]] = []

    examinations = sections.get("examinationIndex") or []
    if examinations:
        for page in paginate_lines(examination_index_lines(sections.get("witnessName"), examinations)):
            pages.append(page)

    exhibits = sections.get("exhibitIndex") or []
    if exhibits:
        for page in paginate_lines(exhibit_index_lines(exhibits)):
            pages.append(page)

    errata = sections.get("errata") or []
    pages.append(changes_page_lines(sections.get("witnessName"), sections.get("depoDate"), errata))
    if len(errata) > CHANGE_ROWS_PAGE1:
        overflow = [
            "  CHANGES AND SIGNATURE (continued)",
            "",
            f"  {'PAGE':<8}{'LINE':<8}{'CHANGE':<30}REASON",
            "",
        ]
        for row in errata[CHANGE_ROWS_PAGE1:]:
            change_text = f"{row.get('from', '')} -> {row.get('to', '')}"[:29]
            overflow.append(
                f"  {_page_str(row.get('page')):<8}{_page_str(row.get('line')):<8}{change_text:<30}{row.get('reason', '')}"
            )
        pages.extend(paginate_lines(overflow))
    pages.append(
        signature_page_lines(
            sections.get("witnessName"),
            notary_county=sections.get("notaryCounty"),
            notary_name=sections.get("notaryName"),
            identification_method=sections.get("identificationMethod"),
        )
    )
    return pages


def write_certified_back_matter(doc: Document, sections: Mapping[str, object]) -> None:
    """Render the certified back-matter pages into an existing document, each as a
    lined page, with a page break between pages."""
    from docx.enum.text import WD_BREAK

    pages = build_certified_back_matter_lines(sections)
    for index, page_lines in enumerate(pages):
        write_lined_page(doc, page_lines)
        if index < len(pages) - 1:
            break_para = doc.add_paragraph()
            break_para.add_run().add_break(WD_BREAK.PAGE)


def render_certified_back_matter_to_docx(sections: Mapping[str, object], output_path: str) -> str:
    """Standalone: render the certified back-matter to its own DOCX (observability +
    local testing without the full document or any deployment)."""
    doc = Document()
    write_certified_back_matter(doc, sections)
    destination = Path(output_path)
    if destination.suffix.lower() != ".docx":
        destination = destination.with_suffix(".docx")
    doc.save(destination)
    return str(destination)


# The deterministic certified section order for a complete document. Front matter is
# unnumbered; the numbered body sits between; certificate + changes/signature close it.
CERTIFIED_DOCUMENT_ORDER = (
    "title_page",
    "caption",
    "witness_index",
    "exhibit_index",
    "body",
    "certificate",
    "changes_signature",
)


def build_certified_front_matter_lines(data: Mapping[str, object]) -> list[list[str]]:
    """Front-matter pages, in certified order: title page, caption/appearances, witness
    index, exhibit index. Index pages are included only when they have rows."""
    caption = data.get("caption") or {}
    reporter = data.get("reporterCertificate") or {}
    pages: list[list[str]] = []
    pages.extend(paginate_lines(title_page_lines(caption, reporter)))
    pages.extend(paginate_lines(caption_lines(caption, data.get("appearances") or [])))

    examinations = data.get("examinationIndex") or []
    if examinations:
        pages.extend(paginate_lines(examination_index_lines(data.get("witnessName") or caption.get("deponent"), examinations)))
    exhibits = data.get("exhibitIndex") or []
    if exhibits:
        pages.extend(paginate_lines(exhibit_index_lines(exhibits)))
    return pages


def build_certificate_and_signature_lines(data: Mapping[str, object]) -> list[list[str]]:
    """Back-matter pages, in certified order: reporter certificate (REQUIRED — raises if
    reporter identity/CSR is absent), then the changes grid + signature/notary page."""
    reporter = data.get("reporterCertificate") or {}
    pages: list[list[str]] = []
    pages.extend(
        paginate_lines(
            certificate_lines(
                reporter,
                reporter_firm=data.get("reporterFirm"),
                reporter_address=data.get("reporterAddress"),
                reporter_phone=data.get("reporterPhone"),
                waiver=data.get("signatureWaived", True),
            )
        )
    )

    witness_name = data.get("witnessName") or (data.get("caption") or {}).get("deponent")
    errata = data.get("errata") or []
    pages.append(changes_page_lines(witness_name, data.get("depoDate"), errata))
    if len(errata) > CHANGE_ROWS_PAGE1:
        overflow = [
            "  CHANGES AND SIGNATURE (continued)",
            "",
            f"  {'PAGE':<8}{'LINE':<8}{'CHANGE':<30}REASON",
            "",
        ]
        for row in errata[CHANGE_ROWS_PAGE1:]:
            change_text = f"{row.get('from', '')} -> {row.get('to', '')}"[:29]
            overflow.append(
                f"  {_page_str(row.get('page')):<8}{_page_str(row.get('line')):<8}{change_text:<30}{row.get('reason', '')}"
            )
        pages.extend(paginate_lines(overflow))
    pages.append(
        signature_page_lines(
            witness_name,
            notary_county=data.get("notaryCounty"),
            notary_name=data.get("notaryName"),
            identification_method=data.get("identificationMethod"),
        )
    )
    return pages


def _write_pages(doc: Document, pages: Sequence[Sequence[str]], *, trailing_break: bool) -> None:
    from docx.enum.text import WD_BREAK

    for index, page_lines in enumerate(pages):
        write_lined_page(doc, page_lines)
        is_last = index == len(pages) - 1
        if not is_last or trailing_break:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def render_certified_document_to_docx(
    render_model: Mapping[str, object],
    certified_data: Mapping[str, object],
    output_path: str,
) -> str:
    """Assemble ONE complete certified DOCX: front matter -> numbered body -> certificate
    + changes/signature. The body is rendered by the surviving body renderer
    (write_render_model_body) — no second transcript-rendering authority. Raises
    CertifiedMetadataError if required reporter certification metadata is absent."""
    from .docx_exporter import apply_body_geometry, write_render_model_body

    doc = Document()
    apply_body_geometry(doc, render_model["geometry"])

    _write_pages(doc, build_certified_front_matter_lines(certified_data), trailing_break=True)
    write_render_model_body(doc, render_model)

    from docx.enum.text import WD_BREAK

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _write_pages(doc, build_certificate_and_signature_lines(certified_data), trailing_break=False)

    destination = Path(output_path)
    if destination.suffix.lower() != ".docx":
        destination = destination.with_suffix(".docx")
    doc.save(destination)
    return str(destination)
